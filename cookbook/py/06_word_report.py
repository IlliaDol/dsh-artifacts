"""A Word report: heading, paragraph, table and an embedded figure made by matplotlib."""
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches

fig, ax = plt.subplots(figsize=(5, 2.8)); x = np.arange(1, 13)
ax.plot(x, 50 + 4 * x + 6 * np.sin(x / 2), color="#2a78d6", lw=2); ax.set(title="Monthly trend"); fig.tight_layout(); fig.savefig("py_out/_fig.png", dpi=150)
d = Document(); d.add_heading("Quarterly report", 0)
d.add_paragraph("Sales grew steadily through the year; the table gives the details.")
t = d.add_table(rows=1, cols=3); t.style = "Light Grid Accent 1"
for c, h in zip(t.rows[0].cells, ["Region", "Q1", "Q2"]): c.text = h
for r in [("North", 120, 150), ("South", 90, 110)]:
    cells = t.add_row().cells
    for c, v in zip(cells, r): c.text = str(v)
d.add_picture("py_out/_fig.png", width=Inches(5)); d.save("py_out/report.docx"); print("written: report.docx")
