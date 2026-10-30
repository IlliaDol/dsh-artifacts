"""An .xlsx with data AND native Excel charts (editable in Excel, not pictures)."""
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference

wb = Workbook(); ws = wb.active; ws.title = "Sales"
ws.append(["Month", "North", "South"])
for m, (n, s) in enumerate([(120, 90), (135, 95), (150, 110), (170, 118), (160, 130), (185, 140)], start=1):
    ws.append([f"M{m}", n, s])
data = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=7)    # includes header row for series names
cats = Reference(ws, min_col=1, min_row=2, max_row=7)
bar = BarChart(); bar.title = "Sales by region"; bar.add_data(data, titles_from_data=True); bar.set_categories(cats)
bar.y_axis.title = "units"; bar.height, bar.width = 7.5, 14
ws.add_chart(bar, "E2")
line = LineChart(); line.title = "Trend"; line.add_data(data, titles_from_data=True); line.set_categories(cats)
line.height, line.width = 7.5, 14
ws.add_chart(line, "E18")
wb.save("py_out/sales_with_charts.xlsx"); print("written: sales_with_charts.xlsx")
