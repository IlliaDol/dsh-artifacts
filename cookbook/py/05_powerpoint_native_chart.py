"""A PowerPoint slide with a title and a NATIVE chart (editable data) plus a text box."""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[5])           # title only
slide.shapes.title.text = "Languages used in the course"
cd = CategoryChartData(); cd.categories = ["Python", "SQL", "R", "Java", "Julia"]
cd.add_series("Users", (68, 57, 41, 29, 12))
gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(1.6), Inches(8), Inches(5), cd)
ch = gf.chart; ch.has_legend = False; ch.value_axis.minimum_scale = 0
box = slide.shapes.add_textbox(Inches(9), Inches(2), Inches(3.8), Inches(3)).text_frame
box.word_wrap = True; box.text = "Python leads, but SQL is close."; box.paragraphs[0].font.size = Pt(22)
prs.save("py_out/chart_slide.pptx"); print("written: chart_slide.pptx")
